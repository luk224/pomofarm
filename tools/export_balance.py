"""Vuelca los valores de pomofarm_balance.xlsx a JSON para los tests de Go.

Uso (desde la raíz):  python3 -I tools/export_balance.py
Escribe backend/internal/game/testdata/balance_xlsx.json. El xlsx es la fuente
de verdad de las cifras; internal/game/config.go debe coincidir con este JSON.
"""
import json
import openpyxl

wb = openpyxl.load_workbook("pomofarm_balance.xlsx", data_only=True)

params = {r[0]: r[1] for r in wb["Parametros"].iter_rows(values_only=True) if r[0] and isinstance(r[1], (int, float))}

crops, flow = [], []
for r in wb["Cultivos"].iter_rows(min_row=4, values_only=True):
    if not r[0] or not isinstance(r[1], (int, float)):
        continue
    row = dict(name=r[0], duration_min=r[1], unlock=r[2], reward=r[3], life_h=r[4],
               e=r[5], yield_coins=r[6], rate_per_h=r[7])
    (flow if r[0] == "Roble Flow" else crops).append(row)

plots, silo, structs = {}, [], {}
for r in wb["Costes"].iter_rows(min_row=4, values_only=True):
    if isinstance(r[0], int):
        plots[r[0]] = r[1]
    if isinstance(r[4], int):
        silo.append(dict(level=r[4], capacity_h=r[5], cost=r[6]))
    elif isinstance(r[4], str) and isinstance(r[5], (int, float)):
        structs[r[4]] = r[5]

out = dict(params=params, crops=crops, flow=flow, plot_costs=plots, silo=silo, structures=structs)
with open("backend/internal/game/testdata/balance_xlsx.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1, sort_keys=True)
print("ok:", len(crops), "cultivos,", len(flow), "flow,", len(plots), "parcelas,", len(silo), "silo")
